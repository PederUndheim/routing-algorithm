r"""Assignments in, results out: the admin side of the validation app.

    $py = "C:\Program Files\QGIS 4.2.0\bin\python-qgis.bat"
    & $py -m validation_app.admin assign                    # upload assignments.json
    & $py -m validation_app.admin export                    # -> data/validation/export/*.xlsx
    & $py -m validation_app.admin backup                    # -> data/validation/backup/*.json

Each talks to the API: --api (default: VALIDATION_API, else localhost) with
--key (default: VALIDATION_ADMIN_KEY, else "dev", the local server's key).

## assignments.json

    {
      "havard":  {"regions": ["Lyngen", "Tromsø"], "ids": [412], "exclude": [101]},
      "markus":  "all",
      "andreas": {"ids": [1, 2, 3]}
    }

Keys are user ids (see backend/verdicts.py). Regions are matched by name
against the corridor index; an unknown name is an error, not an empty set,
so a typo cannot silently hand somebody nothing. A user missing from the
file, or set to "all", sees every corridor. The upload replaces the whole
assignment, so the file is the truth: edit it and run `assign` again.

## The export

One row per corridor. Per reviewer: class (e.g. "blå+"), size, usikker,
problematisk, note. Then what the meeting needs to sort by: how many have
judged it, whether they disagree on class, and the spread in steps, where
one step is one -/0/+ notch and a whole class is three. The test user is
never in it. A second sheet holds every write ever made, oldest first.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from validation_app.backend.verdicts import KAST_CLASSES, TEST_USER, USERS

HERE = Path(__file__).resolve().parent
INDEX = HERE / "frontend" / "public" / "corridors" / "index.json"
ASSIGNMENTS = HERE / "assignments.json"
OUT = HERE.parent / "data" / "validation"

REVIEWERS = [u for u in USERS if u["id"] != TEST_USER]
CLASS_NO = {"green": "grønn", "blue": "blå", "red": "rød", "black": "svart"}
SUB_MARK = {-1: "−", 0: "", 1: "+"}
SIZE_NO = {-1: "mindre", 0: "", 1: "større"}


def _call(args: argparse.Namespace, method: str, path: str, body: Any = None) -> Any:
    request = Request(args.api.rstrip("/") + path, method=method,
                      data=None if body is None else json.dumps(body).encode(),
                      headers={"Content-Type": "application/json", "X-Admin-Key": args.key})
    try:
        with urlopen(request, timeout=120) as response:
            return json.loads(response.read())
    except HTTPError as e:
        raise SystemExit(f"{method} {path}: {e.code} {e.read().decode(errors='replace')}")


def _index() -> list[dict]:
    if not INDEX.exists():
        raise SystemExit(f"No {INDEX} - run `python -m validation_app.prepare` first.")
    return json.loads(INDEX.read_text(encoding="utf-8"))


# --- assign ----------------------------------------------------------------

def resolve(spec: dict[str, Any], index: list[dict]) -> dict[str, list[int]]:
    """assignments.json as explicit id lists; users set to "all" are left out."""
    by_region: dict[str, list[int]] = {}
    for c in index:
        by_region.setdefault(c["region"], []).append(c["id"])
    known_ids = {c["id"] for c in index}
    known_users = {u["id"] for u in USERS}

    resolved = {}
    for user, rule in spec.items():
        if user not in known_users:
            raise SystemExit(f"Unknown user {user!r}; users are {sorted(known_users)}")
        if rule == "all":
            continue
        if not isinstance(rule, dict):
            raise SystemExit(f"{user}: expected \"all\" or an object with regions/ids/exclude")
        ids: set[int] = set()
        for region in rule.get("regions", []):
            if region not in by_region:
                raise SystemExit(f"{user}: unknown region {region!r}. "
                                 f"Regions: {', '.join(by_region)}")
            ids.update(by_region[region])
        ids.update(rule.get("ids", []))
        ids.difference_update(rule.get("exclude", []))
        unknown = ids - known_ids
        if unknown:
            raise SystemExit(f"{user}: no corridor with id {sorted(unknown)}")
        resolved[user] = sorted(ids)
    return resolved


def assign(args: argparse.Namespace) -> None:
    # utf-8-sig: Notepad and PowerShell both like to start a file with a BOM.
    spec = json.loads(Path(args.file).read_text(encoding="utf-8-sig"))
    resolved = resolve(spec, _index())
    for user in USERS:
        n = len(resolved[user["id"]]) if user["id"] in resolved else "alle"
        print(f"  {user['name']:10s} {n}")
    result = _call(args, "PUT", "/api/admin/assignments", {"assignments": resolved})
    print(f"Uploaded: {result}")


# --- export ----------------------------------------------------------------

def _label(verdict: dict[str, Any]) -> str:
    if not verdict.get("kast"):
        return ""
    return CLASS_NO[verdict["kast"]] + SUB_MARK[verdict.get("sub", 0)]


def _step(verdict: dict[str, Any]) -> int:
    return KAST_CLASSES.index(verdict["kast"]) * 3 + verdict.get("sub", 0)


def export(args: argparse.Namespace) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    dump = _call(args, "GET", "/api/admin/dump")
    index = _index()
    current: dict[tuple[str, int], dict] = {
        (v["user"], v["corridor"]): v for v in dump["current"] if v["user"] != TEST_USER}
    assignments = dump["assignments"]

    def assigned(user: str, cid: int) -> bool:
        return user not in assignments or cid in assignments[user]

    book = Workbook()
    sheet = book.active
    sheet.title = "Vurderinger"
    header = ["ID", "Region", "Lengde (m)"]
    for u in REVIEWERS:
        header += [f"{u['name']} KAST", f"{u['name']} størrelse", f"{u['name']} usikker",
                   f"{u['name']} problematisk", f"{u['name']} notat"]
    header += ["Tildelt", "Vurdert", "Ulik klasse", "Spenn (trinn)",
               "Ulik størrelse", "Flagget av noen"]
    sheet.append(header)

    for c in index:
        row: list[Any] = [c["id"], c["region"], c["lengthM"]]
        judged = []
        sizes, flagged, n_assigned = set(), False, 0
        for u in REVIEWERS:
            n_assigned += assigned(u["id"], c["id"])
            v = current.get((u["id"], c["id"]), {})
            if v.get("kast"):
                judged.append(v)
                sizes.add(v.get("size", 0))
            flagged |= bool(v.get("doubt") or v.get("problematic"))
            row += [_label(v), SIZE_NO[v.get("size", 0)],
                    "ja" if v.get("doubt") else "", "ja" if v.get("problematic") else "",
                    v.get("note", "")]
        steps = [_step(v) for v in judged]
        row += [n_assigned, len(judged),
                "ja" if len({v["kast"] for v in judged}) > 1 else "",
                max(steps) - min(steps) if steps else "",
                "ja" if len(sizes) > 1 else "",
                "ja" if flagged else ""]
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    for i, title in enumerate(header, 1):
        sheet.column_dimensions[get_column_letter(i)].width = (
            40 if title.endswith("notat") else max(8, min(len(title) + 2, 18)))

    log = book.create_sheet("Historikk")
    names = {u["id"]: u["name"] for u in USERS}
    log.append(["Tid (server)", "Tid (klient)", "Bruker", "ID", "Endring"])
    for h in dump["history"]:
        if h["user"] != TEST_USER:
            log.append([h["serverTs"], h["clientTs"], names.get(h["user"], h["user"]),
                        h["corridor"], json.dumps(h["patch"], ensure_ascii=False)])
    for cell in log[1]:
        cell.font = Font(bold=True)

    out = Path(args.out) if args.out else (
        OUT / "export" / f"validering_{datetime.now():%Y%m%d_%H%M}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)
    print(f"{len(index)} corridors, {len(current)} verdicts -> {out}")


# --- backup ----------------------------------------------------------------

def backup(args: argparse.Namespace) -> None:
    dump = _call(args, "GET", "/api/admin/dump")
    out = Path(args.out) if args.out else (
        OUT / "backup" / f"validering_{datetime.now():%Y%m%d_%H%M%S}.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(dump, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(dump['current'])} verdicts, {len(dump['history'])} history rows -> {out}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m validation_app.admin",
                                     description=__doc__.split("\n\n")[0])
    parser.add_argument("--api", default=os.environ.get("VALIDATION_API", "http://127.0.0.1:8790"))
    parser.add_argument("--key", default=os.environ.get("VALIDATION_ADMIN_KEY", "dev"))
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("assign", help="upload assignments.json")
    p.add_argument("file", nargs="?", default=ASSIGNMENTS)
    p.set_defaults(func=assign)
    p = sub.add_parser("export", help="the spreadsheet for the consensus meeting")
    p.add_argument("--out")
    p.set_defaults(func=export)
    p = sub.add_parser("backup", help="everything in the store, as JSON")
    p.add_argument("--out")
    p.set_defaults(func=backup)
    args = parser.parse_args(argv)
    args.func(args)
    return 0


if __name__ == "__main__":
    sys.exit(main())
